"""Build the Word report. Run: python3 build_docx.py

Two passes: the document is rendered to PDF with LibreOffice to find the page
of each heading, and those page numbers are written into the (still
updatable) table-of-contents field so the TOC is correct before Word refreshes it.
"""
import os, re, subprocess, shutil, json
from docx import Document
from docx.shared import Pt, Cm, RGBColor, Emu
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_ORIENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import data as D

HERE = os.path.dirname(os.path.abspath(__file__))
OUTDIR = os.path.dirname(HERE)
OUT = os.path.join(OUTDIR, D.REPORT_NAME + ".docx")
CH = os.path.join(HERE, "charts")

NAVY = RGBColor(0x1F, 0x2A, 0x44)
ACCENT = RGBColor(0x2A, 0x78, 0xD6)
GREY = RGBColor(0x52, 0x51, 0x4E)
MUTED = RGBColor(0x8A, 0x89, 0x84)
CONF_COLORS = {"medium": "1F5FAF", "early": "A8501F", "caveat": "52514E"}
FONT = "Calibri"


# ---------- low-level helpers ----------
def shade(cell, hex_fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), hex_fill)
    tcPr.append(shd)


def cell_borders(cell, **kw):
    tcPr = cell._tc.get_or_add_tcPr()
    b = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        spec = kw.get(edge, {"val": "nil"})
        el = OxmlElement(f"w:{edge}")
        for k, v in spec.items():
            el.set(qn(f"w:{k}"), str(v))
        b.append(el)
    tcPr.append(b)


def cell_margins(cell, top=80, bottom=80, left=120, right=120):
    tcPr = cell._tc.get_or_add_tcPr()
    m = OxmlElement("w:tcMar")
    for k, v in (("top", top), ("bottom", bottom), ("left", left), ("right", right)):
        el = OxmlElement(f"w:{k}"); el.set(qn("w:w"), str(v)); el.set(qn("w:type"), "dxa"); m.append(el)
    tcPr.append(m)


def table_borders(table, color="D9D8D3"):
    tblPr = table._tbl.tblPr
    b = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        if edge in ("left", "right", "insideV"):
            el.set(qn("w:val"), "nil")
        else:
            el.set(qn("w:val"), "single"); el.set(qn("w:sz"), "4"); el.set(qn("w:color"), color)
        b.append(el)
    tblPr.append(b)


def repeat_header(row):
    trPr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader"); el.set(qn("w:val"), "true"); trPr.append(el)


def no_split(row):
    trPr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:cantSplit"); el.set(qn("w:val"), "true"); trPr.append(el)


def add_field(paragraph, instr, placeholder="1", bold=False, size=None, color=None):
    r1 = paragraph.add_run(); f1 = OxmlElement("w:fldChar"); f1.set(qn("w:fldCharType"), "begin"); r1._r.append(f1)
    r2 = paragraph.add_run(); t = OxmlElement("w:instrText"); t.set(qn("xml:space"), "preserve"); t.text = f" {instr} "; r2._r.append(t)
    r3 = paragraph.add_run(); f3 = OxmlElement("w:fldChar"); f3.set(qn("w:fldCharType"), "separate"); r3._r.append(f3)
    r4 = paragraph.add_run(placeholder)
    r5 = paragraph.add_run(); f5 = OxmlElement("w:fldChar"); f5.set(qn("w:fldCharType"), "end"); r5._r.append(f5)
    for r in (r4,):
        r.bold = bold
        if size: r.font.size = size
        if color: r.font.color.rgb = color
    return r4


def keep_with_next(p):
    p.paragraph_format.keep_with_next = True


def conf_key(conf):
    c = conf.lower()
    if c.startswith("caveat"): return "caveat"
    if c.startswith("early"): return "early"
    return "medium"


def add_tag(paragraph, text, size=8.5):
    r = paragraph.add_run(f"  [{text.upper()}]")
    r.bold = True; r.font.size = Pt(size)
    r.font.color.rgb = RGBColor.from_string(CONF_COLORS[conf_key(text)])
    return r


# ---------- document ----------
class Builder:
    def __init__(self, toc_pages=None):
        self.doc = Document()
        self.toc_pages = toc_pages or {}
        self.headings = []  # (level, text)
        self.fig_n = 0
        self.tab_n = 0
        self._styles()
        sec = self.doc.sections[0]
        sec.page_width, sec.page_height = Cm(21), Cm(29.7)
        sec.left_margin = sec.right_margin = Cm(2.2)
        sec.top_margin = Cm(2.0); sec.bottom_margin = Cm(2.0)
        sec.different_first_page_header_footer = True
        self._footer(sec)

    @property
    def width(self):
        s = self.doc.sections[-1]
        return s.page_width - s.left_margin - s.right_margin

    def _styles(self):
        st = self.doc.styles
        n = st["Normal"]; n.font.name = FONT; n.font.size = Pt(10.5); n.font.color.rgb = RGBColor(0x2B, 0x2B, 0x2B)
        n.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        n.paragraph_format.space_after = Pt(6); n.paragraph_format.line_spacing = 1.12
        for name, size, before, after in (("Heading 1", 18, 18, 8), ("Heading 2", 13, 12, 4), ("Heading 3", 11, 10, 3)):
            h = st[name]; h.font.name = FONT; h.font.size = Pt(size); h.font.bold = True
            h.font.color.rgb = NAVY if name != "Heading 3" else ACCENT
            rpr = h.element.get_or_add_rPr(); rf = rpr.find(qn("w:rFonts"))
            if rf is None:
                rf = OxmlElement("w:rFonts"); rpr.append(rf)
            for a in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
                rf.set(qn(a), FONT)
            for a in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
                if rf.get(qn(a)) is not None: del rf.attrib[qn(a)]
            h.paragraph_format.space_before = Pt(before); h.paragraph_format.space_after = Pt(after)
            h.paragraph_format.keep_with_next = True
        c = st["Caption"]; c.font.name = FONT; c.font.size = Pt(9); c.font.italic = False
        c.font.color.rgb = GREY; c.font.bold = False
        c.paragraph_format.space_before = Pt(2); c.paragraph_format.space_after = Pt(12)
        for lvl in (1, 2):
            try:
                t = st[f"TOC {lvl}"]
            except KeyError:
                t = st.add_style(f"TOC {lvl}", 1)
            t.font.name = FONT; t.font.size = Pt(11 if lvl == 1 else 10)
            t.paragraph_format.left_indent = Cm(0 if lvl == 1 else 0.6)
            t.paragraph_format.space_after = Pt(3 if lvl == 1 else 1)
            t.paragraph_format.tab_stops.add_tab_stop(self_width_tab(), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
            if lvl == 1: t.font.bold = True

    def _footer(self, sec):
        p = sec.footer.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        r = p.add_run("Supplement brands: cross-brand patterns  |  Page "); r.font.size = Pt(8.5); r.font.color.rgb = MUTED
        add_field(p, "PAGE", "1", size=Pt(8.5), color=MUTED)
        r = p.add_run(" of "); r.font.size = Pt(8.5); r.font.color.rgb = MUTED
        add_field(p, "NUMPAGES", "1", size=Pt(8.5), color=MUTED)

    # --- blocks
    def h(self, text, level=1, page_break=False):
        if page_break:
            self.doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
            # remove the empty paragraph spacing
            self.doc.paragraphs[-1].paragraph_format.space_after = Pt(0)
        p = self.doc.add_heading(text, level=level)
        if level <= 2:
            self.headings.append((level, text))
        return p

    def lead(self, text):
        p = self.doc.add_paragraph()
        r = p.add_run(text); r.bold = True; r.font.size = Pt(11.5); r.font.color.rgb = NAVY
        p.paragraph_format.space_after = Pt(8)
        return p

    def para(self, text, size=None, color=None, italic=False, bold_prefix=None, after=None):
        p = self.doc.add_paragraph()
        if bold_prefix:
            r = p.add_run(bold_prefix); r.bold = True
            if size: r.font.size = Pt(size)
        r = p.add_run(text); r.italic = italic
        if size: r.font.size = Pt(size)
        if color: r.font.color.rgb = color
        if after is not None: p.paragraph_format.space_after = Pt(after)
        return p

    def bullet(self, text, bold_prefix=None, size=None):
        p = self.doc.add_paragraph(style="List Bullet")
        if bold_prefix:
            r = p.add_run(bold_prefix); r.bold = True
            if size: r.font.size = Pt(size)
        r = p.add_run(text)
        if size: r.font.size = Pt(size)
        p.paragraph_format.space_after = Pt(3)
        return p

    def caption(self, kind, text):
        if kind == "Figure":
            self.fig_n += 1; n = self.fig_n
        else:
            self.tab_n += 1; n = self.tab_n
        p = self.doc.add_paragraph(style="Caption")
        r = p.add_run(f"{kind} "); r.bold = True
        add_field(p, f"SEQ {kind} \\* ARABIC", str(n), bold=True)
        r = p.add_run(": "); r.bold = True
        p.add_run(text)
        return p

    def figure(self, png, caption, width_cm=16.4):
        p = self.doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0); keep_with_next(p)
        p.add_run().add_picture(os.path.join(CH, png), width=Cm(width_cm))
        self.caption("Figure", caption)

    def callout(self, title, body, stat=None, stat_label=None, tag=None, fill="EEF4FC", bar="2A78D6"):
        t = self.doc.add_table(rows=1, cols=2 if stat else 1)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = False
        cells = t.rows[0].cells
        no_split(t.rows[0])
        if stat:
            cells[0].width = Cm(3.4); cells[1].width = self.width - Cm(3.4)
            c0 = cells[0]; shade(c0, fill); cell_margins(c0, 70, 70, 160, 60)
            cell_borders(c0, left={"val": "single", "sz": 36, "color": bar})
            c0.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = c0.paragraphs[0]; r = p.add_run(stat); r.bold = True; r.font.size = Pt(20); r.font.color.rgb = ACCENT
            p.paragraph_format.space_after = Pt(0)
            p2 = c0.add_paragraph(); r = p2.add_run(stat_label); r.font.size = Pt(8); r.font.color.rgb = GREY
            p2.paragraph_format.space_after = Pt(0); p2.paragraph_format.line_spacing = 1.0
            c = cells[1]; shade(c, fill); cell_margins(c, 70, 70, 120, 160); cell_borders(c)
        else:
            c = cells[0]; c.width = self.width; shade(c, fill); cell_margins(c, 120, 120, 200, 200)
            cell_borders(c, left={"val": "single", "sz": 36, "color": bar})
        p = c.paragraphs[0]
        r = p.add_run(title); r.bold = True; r.font.size = Pt(11); r.font.color.rgb = NAVY
        if tag: add_tag(p, tag)
        p.paragraph_format.space_after = Pt(3)
        p2 = c.add_paragraph(); r = p2.add_run(body); r.font.size = Pt(9)
        p2.paragraph_format.space_after = Pt(0)
        set_grid(t, [c.width for c in cells])
        sp = self.doc.add_paragraph(); sp.paragraph_format.space_after = Pt(0)
        sp.paragraph_format.line_spacing = 0.6
        return t

    def table(self, header, rows, widths_cm, size=8.5, header_fill="1F2A44", zebra=True, first_bold=True):
        t = self.doc.add_table(rows=1, cols=len(header))
        t.alignment = WD_TABLE_ALIGNMENT.CENTER; t.autofit = False
        table_borders(t)
        hdr = t.rows[0]; repeat_header(hdr); no_split(hdr)
        for i, txt in enumerate(header):
            c = hdr.cells[i]; c.width = Cm(widths_cm[i]); shade(c, header_fill); cell_margins(c, 60, 60, 90, 90)
            p = c.paragraphs[0]; r = p.add_run(txt); r.bold = True; r.font.size = Pt(size); r.font.color.rgb = RGBColor(255, 255, 255)
            p.paragraph_format.space_after = Pt(0)
        for ri, row in enumerate(rows):
            cells = t.add_row().cells
            no_split(t.rows[-1])
            for i, val in enumerate(row):
                c = cells[i]; c.width = Cm(widths_cm[i]); cell_margins(c, 50, 50, 90, 90)
                if zebra and ri % 2 == 1: shade(c, "F6F6F4")
                fill = None
                if isinstance(val, dict):
                    fill = val.get("fill"); txt = val["text"]; bold = val.get("bold", False); colr = val.get("color")
                    align = val.get("align")
                else:
                    txt = str(val); bold = first_bold and i == 0; colr = None; align = None
                if fill: shade(c, fill)
                p = c.paragraphs[0]; r = p.add_run(txt); r.font.size = Pt(size); r.bold = bold
                if colr: r.font.color.rgb = RGBColor.from_string(colr)
                if align == "center": p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_after = Pt(0); p.paragraph_format.line_spacing = 1.05
                c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_grid(t, [Cm(w) for w in widths_cm])
        return t

    def toc(self):
        p = self.doc.add_paragraph()
        r = p.add_run("Contents"); r.bold = True; r.font.size = Pt(18); r.font.color.rgb = NAVY
        p.paragraph_format.space_after = Pt(12)
        # TOC field; the result is pre-filled from the headings list of the previous pass
        p = self.doc.add_paragraph()
        r = p.add_run(); f = OxmlElement("w:fldChar"); f.set(qn("w:fldCharType"), "begin"); f.set(qn("w:dirty"), "true"); r._r.append(f)
        r = p.add_run(); t = OxmlElement("w:instrText"); t.set(qn("xml:space"), "preserve"); t.text = ' TOC \\o "1-2" \\h \\z \\u '; r._r.append(t)
        r = p.add_run(); f = OxmlElement("w:fldChar"); f.set(qn("w:fldCharType"), "separate"); r._r.append(f)
        entries = self.toc_pages.get("entries", [])
        first = True
        for lvl, text, page in entries:
            if first:
                q = p; first = False
                q.style = self.doc.styles[f"TOC {lvl}"]
            else:
                q = self.doc.add_paragraph(style=f"TOC {lvl}")
            q.add_run(f"{text}\t{page}")
        if first:
            p.add_run("Right-click and choose Update Field to build the table of contents.")
            q = p
        r = q.add_run(); f = OxmlElement("w:fldChar"); f.set(qn("w:fldCharType"), "end"); r._r.append(f)

    def conf(self, text):
        p = self.doc.add_paragraph()
        r = p.add_run("Confidence:"); r.font.size = Pt(8.5); r.font.color.rgb = GREY
        add_tag(p, text)
        p.paragraph_format.space_after = Pt(3); keep_with_next(p)
        return p

    def page_break(self):
        self.doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def set_grid(t, widths):
    grid = t._tbl.tblGrid
    for gc, w in zip(grid.findall(qn("w:gridCol")), widths):
        gc.set(qn("w:w"), str(int(Emu(w).twips if hasattr(Emu(w), "twips") else int(w) / 635)))
    tblPr = t._tbl.tblPr
    tw = tblPr.find(qn("w:tblW"))
    if tw is None:
        tw = OxmlElement("w:tblW"); tblPr.append(tw)
    tw.set(qn("w:w"), str(sum(int(int(w) / 635) for w in widths))); tw.set(qn("w:type"), "dxa")
    lay = OxmlElement("w:tblLayout"); lay.set(qn("w:type"), "fixed"); tblPr.append(lay)


def self_width_tab():
    return Cm(21 - 4.4)


# ---------- content ----------
def build(toc_pages=None):
    B = Builder(toc_pages)
    doc = B.doc

    # Title page
    for _ in range(5):
        doc.add_paragraph()
    p = doc.add_paragraph(); r = p.add_run("CROSS-BRAND PATTERN REPORT"); r.bold = True; r.font.size = Pt(10); r.font.color.rgb = ACCENT
    p = doc.add_paragraph(); r = p.add_run(D.TITLE); r.bold = True; r.font.size = Pt(30); r.font.color.rgb = NAVY
    p.paragraph_format.line_spacing = 1.0; p.paragraph_format.space_after = Pt(14)
    p = doc.add_paragraph(); r = p.add_run(D.SUBTITLE); r.font.size = Pt(13); r.font.color.rgb = GREY
    p.paragraph_format.space_after = Pt(40)
    p = doc.add_paragraph(); r = p.add_run(
        "Supplements, wellness and nutricosmetics. Home markets: US, UK, Australia. "
        "Lens: what successful brands share, where they differ, and the warning signs."); r.font.size = Pt(10.5); r.font.color.rgb = GREY
    p = doc.add_paragraph(); r = p.add_run(f"Run date: {D.RUN_DATE}   |   Analyst: brand-pattern-finder   |   Design: report-designer")
    r.font.size = Pt(9.5); r.font.color.rgb = MUTED
    for _ in range(6):
        doc.add_paragraph()
    p = doc.add_paragraph(); r = p.add_run("How to read the confidence tags"); r.bold = True; r.font.size = Pt(9.5); r.font.color.rgb = NAVY
    p.paragraph_format.space_after = Pt(2)
    p = doc.add_paragraph(); r = p.add_run(
        "Patterns are rated MEDIUM (several brands support it) or EARLY SIGNAL (one or two brands). No pattern is rated "
        "higher than medium because the evidence is snippet-based and the set has no failed brands. Dated facts carry "
        "the brand reports' tags: CONFIRMED (brand or investor statement, solid data), REPORTED (credible third party), "
        "INFERRED (analyst reasoning), USER-SUPPLIED (from the user; none in this set).")
    r.font.size = Pt(9); r.font.color.rgb = GREY
    B.page_break()

    # Contents
    B.toc()
    B.page_break()

    # 1. Key findings
    B.h("1. Key findings")
    B.lead("Five successful brands share a founder story, an online-first path and a format that shows the promise; "
           "they also share the same two warning signs. All of it is medium-confidence correlation.")
    for kf in D.KEY_FINDINGS:
        fill, bar = ("F4F4F2", "8A8984") if kf["conf"] == "Caveat" else ("EEF4FC", "2A78D6")
        if kf["refs"] == "W1, W2":
            fill, bar = "FDF0EA", "EB6834"
        B.callout(kf["headline"], kf["detail"], stat=kf["stat"], stat_label=kf["stat_label"], tag=kf["conf"], fill=fill, bar=bar)

    # 2. Summary and caveats
    B.h("2. Summary and caveats", page_break=True)
    B.lead("The shared path is founder problem, visible format, online proof, then retail; three different retail "
           "destinations then lead to success.")
    B.para(D.SUMMARY)
    B.h("Caveats that apply to everything in this report", 2)
    for t, body in D.CAVEATS:
        B.callout(t, body, tag="Caveat", fill="F4F4F2", bar="8A8984")
    B.para("Brands compared (5): MaryRuth's (US), Free Soul (UK), Novomins (UK), Vida Glow (Australia), SmartyPants (US). "
           "No Spate export was available in inputs/category/ and Revuze was not used.", size=9.5, color=GREY)

    # 3. Scorecard
    B.h("3. Scorecard", page_break=True)
    B.lead("Every brand scores 4-5 on story clarity, while story-product fit is the lowest or joint-lowest score for four of the five.")
    B.figure("fig_scorecard.png", "Brand scorecards on five criteria (1-5). Source: brand reports, 2026-10-06.")
    rows = []
    for k in D.BRAND_KEYS:
        s = D.SCORES[k]
        row = [D.BRAND_META[k]["name"]]
        for v in s:
            fill = {5: "2A78D6", 4: "9EC5F4", 3.5: "CDE2FB", 3: "E8F1FC"}.get(v, "FFFFFF")
            row.append({"text": f"{v:g}", "fill": fill, "align": "center", "bold": True,
                        "color": "FFFFFF" if v == 5 else "1F2A44"})
        rows.append(row)
    B.caption("Table", "Scorecard as a heat-map table (darker = higher). Novomins' overall score is 3.5.")
    B.doc.paragraphs[-1].paragraph_format.keep_with_next = True
    B.table(["Brand"] + D.CRITERIA, rows, [3.6, 2.5, 2.5, 2.5, 2.5, 2.5], size=9)
    B.para("")
    B.para("Reading the scores: story clarity is the highest-scoring criterion in the set (P1). Story-product fit is pulled "
           "down by range sprawl into trend SKUs (W1). Novomins scores 3 on marketing-channel fit because little social or "
           "creator activity was found, which the report treats as a gap rather than evidence of absence.", size=9.5)

    # 4. Comparison matrix
    B.h("4. How the five brands compare", page_break=True)
    B.lead("All five tell a founder-problem story, but their hero formats, price tiers and success measures differ widely.")
    B.h("Story, hero product and price", 2)
    B.caption("Table", "Story type, hero product and price tier. Source: pattern report comparison matrix.")
    B.doc.paragraphs[-1].paragraph_format.keep_with_next = True
    B.table(["Brand", "Story type", "Hero product and format", "Price tier"],
            [[D.MATRIX[k]["label"], D.MATRIX[k]["story"], D.MATRIX[k]["hero"], D.MATRIX[k]["price"]] for k in D.BRAND_KEYS],
            [2.9, 5.0, 5.0, 3.7])
    B.h("Marketing and sales channels", 2)
    B.caption("Table", "Lead marketing channel, first sales channel and channel expansion sequence. Source: pattern report.")
    B.doc.paragraphs[-1].paragraph_format.keep_with_next = True
    B.table(["Brand", "Lead marketing channel", "First sales channel", "Channel expansion sequence"],
            [[D.MATRIX[k]["label"], D.MATRIX[k]["marketing"], D.MATRIX[k]["first_sales"],
              " > ".join(D.MATRIX[k]["sequence"]) + "; " + D.MATRIX[k]["seq_note"]] for k in D.BRAND_KEYS],
            [2.9, 4.4, 3.2, 6.1])
    B.h("Success measures", 2)
    B.lead("Success is measured differently for each brand, so the figures cannot be ranked against each other.")
    B.caption("Table", "Success measure per brand, with the evidence tag from the brand report.")
    B.doc.paragraphs[-1].paragraph_format.keep_with_next = True
    B.table(["Brand", "Success measure", "Evidence"],
            [[D.MATRIX[k]["label"], D.MATRIX[k]["success"], D.MATRIX[k]["success_tag"]] for k in D.BRAND_KEYS],
            [2.9, 10.3, 3.4])

    # 5. Shared patterns
    B.h("5. What the successful brands share", page_break=True)
    B.lead("Two traits (founder-problem origin, women as buyer) appear in all five brands; four more appear in four; "
           "none is rated above medium confidence.")
    B.figure("fig_pattern_counts.png", "Number of brands supporting each pattern. Source: pattern report, 2026-10-06.")
    B.figure("fig_heatmap.png", "Pattern-by-brand matrix: which brands support each pattern, which are named exceptions, "
             "and where evidence was not found. Source: pattern report, 2026-10-06.")
    for p in [x for x in D.PATTERNS if x["group"] == "Shared"]:
        hp = B.h(f"{p['id']}. {p['name']}", 2)
        B.conf(p["conf"])
        sup = [D.BRAND_META[k]["name"] for k in D.BRAND_KEYS if p["cells"][k] == "S"]
        B.para(p["desc"])
        B.bullet(", ".join(sup) + f" ({p['count']})", bold_prefix="Supporting: ", size=9.5)
        B.bullet(p["exceptions"], bold_prefix="Exceptions and caveats: ", size=9.5)

    # 6. Routes, timeline, channels
    B.h("6. Different routes to the same result", page_break=True)
    B.lead("The online-first start is shared, but the retail destination, the geography and the ownership outcome differ by brand.")
    B.figure("fig_timeline.png", "Dated milestones per brand (online launch, physical retail, social commerce, funding, other). "
             "Source: brand-report timelines, 2026-10-06.", width_cm=16.6)
    yes = "2A78D6"; first = "1F2A44"; no = "F3C4AD"
    rows = []
    for k in D.BRAND_KEYS:
        row = [D.BRAND_META[k]["name"]]
        for v in D.CHANNELS[k]:
            if v == "F": row.append({"text": "First", "fill": first, "color": "FFFFFF", "bold": True, "align": "center"})
            elif v == "Y": row.append({"text": "Yes", "fill": yes, "color": "FFFFFF", "align": "center"})
            elif v == "N": row.append({"text": "No", "fill": no, "color": "1F2A44", "align": "center"})
            else: row.append({"text": "-", "color": "8A8984", "align": "center"})
        rows.append(row)
    B.caption("Table", "Sales-channel presence as listed in the pattern report. First = first sales channel; "
              "- = not listed (not proof of absence); No = stated as absent or unofficial.")
    B.doc.paragraphs[-1].paragraph_format.keep_with_next = True
    B.table(["Brand"] + D.CHANNEL_COLS, rows, [2.6] + [1.75] * 8, size=8, zebra=False)
    B.para("", after=2)
    for k in D.BRAND_KEYS:
        B.bullet(D.CHANNEL_NOTES[k], bold_prefix=D.BRAND_META[k]["name"] + ": ", size=9)
    for r_ in D.ROUTES:
        hp = B.h(f"{r_['id']}. {r_['name']}", 2)
        B.conf(r_["conf"])
        B.para(r_["text"])

    # 7. Warning signs
    B.h("7. Warning signs: where story and product or channel pulled apart", page_break=True)
    B.lead("Range sprawl and claims escalation appear in all five brands; only SmartyPants has faced legal action so far.")
    for p in [x for x in D.PATTERNS if x["group"] == "Warning"]:
        hp = B.h(f"{p['id']}. {p['name']}", 2)
        B.conf(p["conf"])
        B.para(p["desc"])
        sup = [D.BRAND_META[k]["name"] for k in D.BRAND_KEYS if p["cells"][k] == "S"]
        B.bullet(", ".join(sup) + f" ({p['count']})", bold_prefix="Brands affected: ", size=9.5)
        B.bullet(p["exceptions"], bold_prefix="What followed / exceptions: ", size=9.5)

    # 8. Success models
    B.h("8. Distinct success models", page_break=True)
    B.lead("Four models explain the five successes; only the mass engine rests on more than one brand.")
    B.caption("Table", "Success models, the brands that fit and confidence. Source: pattern report.")
    B.doc.paragraphs[-1].paragraph_format.keep_with_next = True
    B.table(["Model", "How it works", "Brands that fit", "Confidence"],
            [[m["name"], m["how"], m["brands"], m["conf"]] for m in D.MODELS], [3.3, 7.0, 3.6, 2.7])
    B.h("Alternative explanations", 2)
    B.para("Timing, category growth, funding, founder networks and uneven measurement each explain part of the result, "
           "so the patterns above are correlations, not proven causes.", bold_prefix="")
    for t, body in D.ALTERNATIVES:
        B.bullet(body, bold_prefix=t + ". ", size=9.5)

    # 9. Implications
    B.h("9. Implications for a new brand or product launch", page_break=True)
    B.lead("A new launch should pick a self-explaining format, prove demand online, choose one retail destination and "
           "write claims for the strictest market from day one.")
    for i, (t, body) in enumerate(D.IMPLICATIONS, 1):
        p = B.para(body, bold_prefix=f"{i}. {t} ", size=10)

    # 10. Open questions
    B.h("10. Open questions and gaps", page_break=True)
    B.lead("Full-page evidence, channel revenue mix and failed comparison brands are the three gaps that would most "
           "change confidence in these patterns.")
    B.h("Evidence quality", 2)
    for g in D.GAPS_EVIDENCE:
        B.bullet(g, size=9.5)
    B.h("Failed or stalled brands to add (the set has none)", 2)
    B.para(D.GAPS_BRANDS_NOTE, size=9.5, italic=True, color=GREY)
    for n, t in D.GAPS_BRANDS:
        B.bullet(t, bold_prefix=n + ": ", size=9.5)
    B.h("Recommended Spate category searches", 2)
    B.para("Adding these exports to inputs/category/ would let the next run separate brand success from category growth.", size=9.5)
    B.caption("Table", "Recommended Spate searches. Source: pattern report.")
    B.doc.paragraphs[-1].paragraph_format.keep_with_next = True
    B.table(["Spate search", "Markets", "What it would test"], [list(x) for x in D.SPATE], [6.4, 2.4, 7.8], size=8.5, first_bold=False)

    # 11. Sources
    B.h("11. Sources", page_break=True)
    B.lead("Every finding comes from the 2026-10-06 pattern report; the brand reports supplied only scorecard scores and milestone dates for charts.")
    for path, desc in D.SOURCES:
        B.bullet(desc, bold_prefix=path + ": ", size=9.5)
    B.para(D.SOURCES_NOTE, size=9.5, color=GREY)
    return B


def pdf_headings(docx_path, headings):
    tmp = os.path.join(HERE, "_pdf")
    shutil.rmtree(tmp, ignore_errors=True); os.makedirs(tmp)
    subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", tmp, docx_path],
                   check=True, capture_output=True, timeout=180)
    pdf = os.path.join(tmp, os.path.basename(docx_path).replace(".docx", ".pdf"))
    txt = subprocess.run(["pdftotext", "-layout", pdf, "-"], capture_output=True, text=True).stdout
    pages = txt.split("\f")
    norm = lambda s: re.sub(r"\s+", " ", s).strip()
    found = []
    start = 2  # skip title and contents pages
    for lvl, text in headings:
        key = norm(text)[:40]
        pg = None
        for i in range(start, len(pages)):
            if key in norm(pages[i]):
                pg = i + 1; start = i; break
        found.append((lvl, text, pg if pg else ""))
    return found, pdf, len(pages)


if __name__ == "__main__":
    B = build()
    B.doc.save(OUT)
    entries, _, _ = pdf_headings(OUT, B.headings)
    B2 = build({"entries": entries})
    B2.doc.save(OUT)
    entries2, pdf, n = pdf_headings(OUT, B2.headings)
    mismatch = [(a, b) for a, b in zip(entries, entries2) if a != b]
    if mismatch:  # TOC length changed the pagination; one more pass
        B3 = build({"entries": entries2}); B3.doc.save(OUT)
        entries2, pdf, n = pdf_headings(OUT, B3.headings)
    shutil.copy(pdf, os.path.join(HERE, D.REPORT_NAME + "-preview.pdf"))
    print("saved", OUT)
    print(json.dumps(entries2, indent=0))
